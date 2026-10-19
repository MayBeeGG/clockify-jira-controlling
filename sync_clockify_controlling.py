#!/usr/bin/env python3
"""Refresh Clockify tracked time and Jira original estimates without storing keys in the workbook."""

from __future__ import annotations

import argparse
import base64
import calendar
import json
import os
import re
import sys
import time
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).resolve().parent
NAVY = "17324F"
BLUE = "2F6FA3"
PALE = "EEF5FB"
LIGHT = "F7FAFD"
TEXT = "263746"


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip("\"'"))


def api_json(base: str, path: str, key: str, payload=None, params=None):
    url = base.rstrip("/") + path
    if params:
        url += "?" + urlencode(params)
    body = json.dumps(payload).encode() if payload is not None else None
    headers = {"X-Api-Key": key, "Accept": "application/json"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    for attempt in range(5):
        try:
            with urlopen(Request(url, data=body, headers=headers,
                                 method="POST" if body is not None else "GET"), timeout=90) as response:
                return json.load(response), response.headers.get("Last-Page")
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:300]
            if exc.code in (429, 500, 502, 503, 504) and attempt < 4:
                time.sleep(min(2 ** attempt, 16))
                continue
            raise RuntimeError(f"Clockify HTTP {exc.code} at {path}: {detail}") from exc
        except URLError as exc:
            raise RuntimeError(f"Clockify connection failed: {exc.reason}") from exc


def pages(base: str, path: str, key: str, page_size=200):
    result = []
    for page in range(1, 10001):
        batch, last = api_json(base, path, key, params={"page": page, "page-size": page_size})
        if not isinstance(batch, list):
            raise RuntimeError(f"Unexpected list response from {path}")
        result.extend(batch)
        if last == "true" or (last is None and len(batch) < page_size):
            return result
    raise RuntimeError(f"Pagination limit reached at {path}")


def jira_issues(base: str, email: str, token: str, project_key: str = ""):
    parsed = urlparse(base)
    if parsed.scheme != "https" or not parsed.netloc or parsed.path not in ("", "/"):
        raise RuntimeError("JIRA_URL must be an HTTPS Jira site URL, e.g. https://example.atlassian.net")
    credentials = base64.b64encode(f"{email}:{token}".encode()).decode("ascii")
    headers = {"Authorization": f"Basic {credentials}", "Accept": "application/json",
               "Content-Type": "application/json"}
    try:
        with urlopen(Request(base.rstrip("/") + "/rest/api/3/myself", headers=headers),
                     timeout=30) as response:
            json.load(response)
    except HTTPError as exc:
        if exc.code == 401:
            raise RuntimeError("Jira authentication failed (HTTP 401): use a personal Atlassian API token for JIRA_EMAIL, not an organization API key") from exc
        raise RuntimeError(f"Jira account check failed (HTTP {exc.code})") from exc
    except URLError as exc:
        raise RuntimeError(f"Jira connection failed: {exc.reason}") from exc
    jql_key = project_key.replace('\\', '\\\\').replace('"', '\\"')
    payload = {"jql": f'project = "{jql_key}"' if project_key else "",
               "fields": ["summary", "timeoriginalestimate"],
               "maxResults": 100}
    issues = []
    seen_tokens = set()
    for _ in range(10000):
        body = json.dumps(payload).encode("utf-8")
        for attempt in range(5):
            try:
                request = Request(base.rstrip("/") + "/rest/api/3/search/jql", data=body,
                                  headers=headers, method="POST")
                with urlopen(request, timeout=90) as response:
                    data = json.load(response)
                break
            except HTTPError as exc:
                detail = exc.read().decode("utf-8", errors="replace")[:300]
                if exc.code in (429, 500, 502, 503, 504) and attempt < 4:
                    time.sleep(min(2 ** attempt, 16))
                    continue
                raise RuntimeError(f"Jira HTTP {exc.code}: {detail}") from exc
            except URLError as exc:
                raise RuntimeError(f"Jira connection failed: {exc.reason}") from exc
        batch = data.get("issues")
        if not isinstance(batch, list):
            raise RuntimeError("Jira search returned no issues list")
        issues.extend(batch)
        next_token = data.get("nextPageToken")
        if data.get("isLast") is True or not next_token:
            return issues
        if next_token in seen_tokens:
            raise RuntimeError("Jira search repeated a page token")
        seen_tokens.add(next_token)
        payload["nextPageToken"] = next_token
    raise RuntimeError("Jira search pagination limit reached")


def jira_matches(issues):
    matches = defaultdict(list)
    for issue in issues:
        fields = issue.get("fields") or {}
        name = " ".join((fields.get("summary") or "").split()).casefold()
        if name:
            matches[name].append(issue)
    return matches


def jira_estimate(task_name: str, matches):
    name = " ".join(task_name.split()).casefold()
    found = matches.get(name, []) if name else []
    if not found:
        return None, "No Jira match"
    if len(found) > 1:
        return None, "Ambiguous: " + ", ".join(sorted(issue["key"] for issue in found))
    issue = found[0]
    seconds = (issue.get("fields") or {}).get("timeoriginalestimate")
    return (float(seconds) / 3600 if seconds is not None else None,
            issue["key"] if seconds is not None else issue["key"] + " (no estimate)")


def month_windows(first: date, last: date):
    current = first
    while current <= last:
        month_end = date(current.year, current.month, calendar.monthrange(current.year, current.month)[1])
        end = min(last, month_end)
        yield current, end
        current = end + timedelta(days=1)


def report_entries(base: str, workspace: str, key: str, first: date, last: date):
    output = {}
    for start, end in month_windows(first, last):
        for page in range(1, 10001):
            payload = {
                "dateRangeStart": f"{start}T00:00:00.000",
                "dateRangeEnd": f"{end}T23:59:59.999",
                "dateRangeType": "ABSOLUTE", "exportType": "JSON",
                "detailedFilter": {"page": page, "pageSize": 1000,
                                   "sortColumn": "ID", "options": {"totals": "CALCULATE"}},
            }
            data, _ = api_json(base, f"/workspaces/{workspace}/reports/detailed", key, payload)
            batch = data.get("timeentries", data.get("timeEntries"))
            if not isinstance(batch, list):
                raise RuntimeError("Clockify detailed report has no timeentries list")
            for item in batch:
                item_id = item.get("id") or item.get("_id")
                if not item_id:
                    raise RuntimeError("A Clockify time entry has no ID")
                output[item_id] = item
            if len(batch) < 1000:
                break
        else:
            raise RuntimeError("Detailed report pagination limit reached")
    return list(output.values())


def duration_hours(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value) / 3600
    if isinstance(value, str) and value.isdigit():
        return float(value) / 3600
    match = re.fullmatch(r"P(?:(\d+)D)?(?:T(?:(\d+)H)?(?:(\d+)M)?(?:(\d+(?:\.\d+)?)S)?)?", str(value))
    if not match:
        raise RuntimeError(f"Unrecognized Clockify duration: {value!r}")
    days, hours, minutes, seconds = (float(part or 0) for part in match.groups())
    return days * 24 + hours + minutes / 60 + seconds / 3600


def entry_hours(item):
    interval = item.get("timeInterval") or {}
    value = interval.get("duration") if isinstance(interval, dict) else None
    if value is None:
        value = item.get("duration")
    if value is None:
        raise RuntimeError(f"Entry {item.get('id')} has no completed duration; stop running timers before refresh")
    return duration_hours(value)


def old_manual(ws):
    saved = {}
    ids_at_right = ws["M7"].value == "Task ID"
    for row in ws.iter_rows(min_row=8, values_only=True):
        task_id = row[12] if ids_at_right and len(row) > 13 else row[9]
        project_id = row[13] if ids_at_right and len(row) > 13 else row[10]
        if task_id or project_id:
            saved[(str(project_id or ""), str(task_id or ""))] = (row[8], row[11])
    return saved


def table_header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 34


def add_bar_chart(ws, title: str, category_col: int, value_col: int,
                  header_row: int, last_row: int, anchor: str) -> None:
    if last_row <= header_row:
        return
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = title
    chart.legend = None
    chart.width = 17
    chart.height = max(7, min(12, 2.5 + (last_row - header_row) * 0.75))
    chart.x_axis.title = "Hours"
    chart.add_data(Reference(ws, min_col=value_col, min_row=header_row,
                             max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=category_col,
                                   min_row=header_row + 1, max_row=last_row))
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].graphicalProperties.line.noFill = True
    ws.add_chart(chart, anchor)


def refresh(workbook_path: Path, projects, tasks, users, entries, jira_issues_list, start: date, end: date):
    if workbook_path.exists():
        wb = load_workbook(workbook_path)
    else:
        wb = Workbook()
        wb.active.title = "Dashboard"
        wb.create_sheet("Time entries")
        dashboard = wb["Dashboard"]
        dashboard["A2"], dashboard["A3"] = "Period", "Refreshed"
        dashboard["A4"], dashboard["C4"], dashboard["E4"] = "Tracked hours", "Estimate hours", "Overrun tasks"
        for column, width in {"A": 28, "B": 42, "D": 18, "E": 18, "F": 18, "G": 24, "I": 25}.items():
            dashboard.column_dimensions[column].width = width
        for column in ("A", "B"):
            wb["Time entries"].column_dimensions[column].width = 25
        for column in ("C", "D"):
            wb["Time entries"].column_dimensions[column].width = 30
    ws, raw = wb["Dashboard"], wb["Time entries"]
    people_ws = wb["People"] if "People" in wb else wb.create_sheet("People", 1)
    if wb.sheetnames.index("People") != 1:
        wb.move_sheet(people_ws, offset=1 - wb.sheetnames.index("People"))
    manual = old_manual(ws)
    project_names = {str(p["id"]): p.get("name") or "Unnamed project" for p in projects}
    task_map = {(str(t.get("projectId") or ""), str(t["id"])): t for t in tasks}
    matches = jira_matches(jira_issues_list)
    user_names = {str(u["id"]): u.get("name") or "Unknown user" for u in users}
    hours = defaultdict(float)
    task_user_hours = defaultdict(lambda: defaultdict(float))
    person_hours = defaultdict(float)
    person_billable = defaultdict(float)
    person_entries = defaultdict(int)
    person_projects = defaultdict(lambda: defaultdict(float))
    person_tasks = defaultdict(float)
    project_hours = defaultdict(float)
    raw_rows = []
    for item in entries:
        pid = str(item.get("projectId") or "")
        tid = str(item.get("taskId") or "")
        uid = str(item.get("userId") or item.get("userName") or "unknown")
        if uid not in user_names:
            user_names[uid] = item.get("userName") or "Unknown user"
        person = user_names[uid]
        amount = entry_hours(item)
        hours[(pid, tid)] += amount
        task_user_hours[(pid, tid)][uid] += amount
        person_hours[uid] += amount
        person_billable[uid] += amount if item.get("billable") else 0
        person_entries[uid] += 1
        person_projects[uid][pid] += amount
        person_tasks[(uid, pid, tid)] += amount
        project_hours[pid] += amount
        interval = item.get("timeInterval") or {}
        project = item.get("projectName") or project_names.get(pid) or "No project"
        task_name = item.get("taskName") or task_map.get((pid, tid), {}).get("name") or "No task"
        raw_rows.append([item.get("id"), interval.get("start") or item.get("start"),
                         interval.get("end") or item.get("end"), project, task_name,
                         person, amount, bool(item.get("billable")),
                         item.get("description") or "", pid, tid])

    rows = []
    for pid, tid in set(task_map) | set(hours):
        task = task_map.get((pid, tid), {})
        priority, override = manual.get((pid, tid), (None, None))
        estimate, jira_match = jira_estimate(task.get("name") or "", matches)
        if override is not None and not isinstance(override, (int, float)):
            raise RuntimeError(f"Estimate override for task {tid} must be a number")
        contributions = task_user_hours[(pid, tid)]
        top_name = ""
        if contributions:
            maximum = max(contributions.values())
            leaders = sorted(user_names[uid] for uid, value in contributions.items()
                             if abs(value - maximum) < 1e-8)
            top_name = ("Tie: " if len(leaders) > 1 else "") + ", ".join(leaders)
        rows.append([project_names.get(pid, "No project"), task.get("name") or "No task",
                     estimate, hours.get((pid, tid), 0), None, task.get("status") or "",
                     None, top_name, priority, tid, pid, override, jira_match])
    rows.sort(key=lambda r: (str(r[0]).lower(), str(r[1]).lower(), r[9]))
    total_hours = sum(person_hours.values())
    people_rows = []
    for uid, name in user_names.items():
        people_rows.append([name, person_hours[uid], person_billable[uid],
                            person_hours[uid] - person_billable[uid], None,
                            person_entries[uid],
                            sum(value > 0 for value in person_projects[uid].values()), uid])
    people_rows.sort(key=lambda r: (-r[1], r[0].lower(), r[7]))
    project_rows = sorted([(project_names.get(pid, "No project"), value) for pid, value in project_hours.items()],
                          key=lambda r: (-r[1], r[0].lower()))
    person_task_rows = sorted([(user_names[uid], project_names.get(pid, "No project"),
                                task_map.get((pid, tid), {}).get("name") or "No task", value)
                               for (uid, pid, tid), value in person_tasks.items()],
                              key=lambda r: (r[0].lower(), -r[3], r[1].lower(), r[2].lower()))

    for sheet, first in ((ws, 8), (people_ws, 8), (raw, 2)):
        if sheet.max_row >= first:
            sheet.delete_rows(first, sheet.max_row - first + 1)
        sheet._charts.clear()
        sheet.sheet_view.showGridLines = False

    ws["A1"] = "Clockify time / Jira estimates"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    ws["B2"] = f"{start.isoformat()} to {end.isoformat()}"
    ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws["G4"] = "Estimates not set"
    ws["I4"] = "Unassigned hours"
    ws["J5"] = None
    last_row = max(8, 7 + len(rows))
    ws["B5"] = f"=SUM(D8:D{last_row})"
    ws["D5"] = f'=SUM(C8:C{last_row})+SUM(L8:L{last_row})-SUMIF(L8:L{last_row},"<>",C8:C{last_row})'
    ws["F5"] = f'=COUNTIF(G8:G{last_row},"Overrun")'
    ws["H5"] = f'=COUNTIF(G8:G{last_row},"Estimate not set")'
    ws["I5"] = f'=SUMIF(G8:G{last_row},"Unassigned time",D8:D{last_row})'
    for col in ("A", "C", "E", "G", "I"):
        ws[f"{col}4"].font = Font(name="Arial", size=10, bold=True, color=NAVY)
        ws[f"{col}4"].fill = PatternFill("solid", fgColor=PALE)
        ws[f"{col}5"].fill = PatternFill("solid", fgColor=PALE)
    for col in ("B", "D", "F", "H", "I"):
        ws[f"{col}5"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
        ws[f"{col}5"].number_format = "0.00" if col in ("B", "D", "I") else "0"
        ws[f"{col}5"].alignment = Alignment(horizontal="right")
    table_header(ws, 7, ["Project", "Task", "Jira original estimate (h)", "Tracked (h)",
                         "Remaining (h)", "Status", "Result", "Top contributor",
                         "Priority / note", "", "", "Estimate override (h)",
                         "Task ID", "Project ID", "Jira match"])
    for idx, values in enumerate(rows, 8):
        for col, value in enumerate(values[:9], 1):
            cell = ws.cell(idx, col, value)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if idx % 2 == 0 else LIGHT)
        ws.cell(idx, 12, values[11])
        ws.cell(idx, 13, values[9])
        ws.cell(idx, 14, values[10])
        ws.cell(idx, 15, values[12])
        ws.row_dimensions[idx].height = 23
        ws.cell(idx, 5, f'=IF(COUNT(L{idx})>0,L{idx}-D{idx},IF(COUNT(C{idx})>0,C{idx}-D{idx},""))')
        ws.cell(idx, 7, f'=IF(OR(M{idx}="",N{idx}=""),"Unassigned time",IF(COUNT(C{idx})+COUNT(L{idx})=0,"Estimate not set",IF(E{idx}<0,"Overrun",IF(F{idx}="DONE","Done",IF(D{idx}=0,"Not started","On track")))))')
        if not values[9] or not values[10]:
            ws.cell(idx, 7).fill = PatternFill("solid", fgColor="FCE8E6")
        if len(values[7]) > 27:
            ws.cell(idx, 8).alignment = Alignment(vertical="center", wrap_text=True)
            ws.row_dimensions[idx].height = 34
        for col in (3, 4, 5, 12):
            ws.cell(idx, col).number_format = "0.00"
        for col in (9, 12):
            ws.cell(idx, col).fill = PatternFill("solid", fgColor="FFF3D8")
            ws.cell(idx, col).font = Font(name="Arial", size=10, color=NAVY)
    ws.column_dimensions["H"].width = 31
    ws.column_dimensions["J"].hidden = True
    ws.column_dimensions["K"].hidden = True
    ws.column_dimensions["M"].hidden = True
    ws.column_dimensions["N"].hidden = True
    ws.column_dimensions["L"].width = 22
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["O"].width = 29
    ws.freeze_panes = "C8"
    ws.auto_filter.ref = f"A7:O{last_row}"
    project_header = last_row + 4
    ws.cell(project_header, 1, "Hours by project").font = Font(name="Arial", size=13, bold=True, color=NAVY)
    table_header(ws, project_header + 2, ["Project", "Tracked (h)"])
    for idx, values in enumerate(project_rows, project_header + 3):
        ws.cell(idx, 1, values[0])
        ws.cell(idx, 2, values[1]).number_format = "0.00"
    add_bar_chart(ws, "Tracked hours by project", 1, 2,
                  project_header + 2, project_header + 2 + len(project_rows),
                  f"D{project_header + 1}")

    people_ws["A1"] = "Time by person"
    people_ws["A1"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    people_ws["B2"] = f"{start.isoformat()} to {end.isoformat()}"
    people_ws["A4"], people_ws["D4"], people_ws["G4"] = "Total hours", "People with time", "Billable hours"
    last_person = max(8, 7 + len(people_rows))
    people_ws["B5"] = f"=SUM(B8:B{last_person})"
    people_ws["E5"] = f'=COUNTIF(B8:B{last_person},">0")'
    people_ws["G5"] = f"=SUM(C8:C{last_person})"
    people_ws["H5"] = None
    for label in ("A4", "D4", "G4"):
        people_ws[label].font = Font(name="Arial", size=10, bold=True, color=NAVY)
        people_ws[label].fill = PatternFill("solid", fgColor=PALE)
    for cell in ("B5", "E5", "G5"):
        people_ws[cell].font = Font(name="Arial", size=14, bold=True, color=NAVY)
        people_ws[cell].number_format = "0" if cell == "E5" else "0.00"
    table_header(people_ws, 7, ["Person", "Tracked (h)", "Billable (h)",
                                "Non-billable (h)", "Share", "Entries",
                                "Projects"])
    people_ws["H7"] = None
    for idx, values in enumerate(people_rows, 8):
        for col, value in enumerate(values[:7], 1):
            cell = people_ws.cell(idx, col, value)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if idx % 2 == 0 else LIGHT)
            cell.alignment = Alignment(vertical="center")
        people_ws.cell(idx, 5, f'=IF($B$5=0,0,B{idx}/$B$5)').number_format = "0.0%"
        people_ws.cell(idx, 7).alignment = Alignment(horizontal="right", vertical="center")
        for col in (2, 3, 4):
            people_ws.cell(idx, col).number_format = "0.00"
        people_ws.row_dimensions[idx].height = 23
    people_ws.column_dimensions["A"].width = 28
    people_ws.column_dimensions["B"].width = 17
    people_ws.column_dimensions["C"].width = 18
    people_ws.column_dimensions["D"].width = 22
    people_ws.column_dimensions["E"].width = 12
    people_ws.column_dimensions["F"].width = 12
    people_ws.column_dimensions["G"].width = 14
    people_ws.column_dimensions["H"].hidden = False
    people_ws.column_dimensions["H"].width = 3
    people_ws.freeze_panes = "B8"
    people_ws.auto_filter.ref = f"A7:G{last_person}"
    add_bar_chart(people_ws, "Tracked hours by person", 1, 2, 7,
                  7 + len(people_rows), "J7")
    detail_row = max(last_person + 4, 19)
    people_ws.cell(detail_row, 1, "Person and task").font = Font(name="Arial", size=13, bold=True, color=NAVY)
    table_header(people_ws, detail_row + 2, ["Person", "Project", "Task", "Tracked (h)"])
    for idx, values in enumerate(person_task_rows, detail_row + 3):
        for col, value in enumerate(values, 1):
            people_ws.cell(idx, col, value)
        people_ws.cell(idx, 4).number_format = "0.00"

    for col in range(9, 16):
        raw.cell(1, col).value = None
    table_header(raw, 1, ["Start", "End", "Project", "Task", "Person",
                          "Hours", "Billable", "Description"])
    raw["M1"], raw["N1"], raw["O1"] = "Entry ID", "Project ID", "Task ID"
    for idx, values in enumerate(raw_rows, 2):
        visible = values[1:9]
        for col, value in enumerate(visible, 1):
            cell = raw.cell(idx, col, value)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if idx % 2 == 0 else LIGHT)
        raw.cell(idx, 6).number_format = "0.00"
        raw.cell(idx, 13, values[0])
        raw.cell(idx, 14, values[9])
        raw.cell(idx, 15, values[10])
    raw.column_dimensions["A"].hidden = False
    for col in ("M", "N", "O"):
        raw.column_dimensions[col].hidden = True
    raw.column_dimensions["E"].width = 26
    raw.column_dimensions["H"].width = 40
    raw.freeze_panes = "C2"
    raw.auto_filter.ref = f"A1:H{max(2, len(raw_rows) + 1)}"
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    temporary = workbook_path.with_name(workbook_path.stem + ".refreshing.xlsx")
    try:
        wb.save(temporary)
        temporary.replace(workbook_path)
    finally:
        temporary.unlink(missing_ok=True)
    return {"tasks": rows, "entries": raw_rows, "people": people_rows,
            "person_tasks": person_task_rows, "projects": project_rows,
            "total_hours": total_hours, "billable_hours": sum(person_billable.values()),
            "estimates_not_set": sum(bool(row[9]) and row[2] is None and row[11] is None for row in rows),
            "unassigned_hours": sum(value for (pid, tid), value in hours.items() if not pid or not tid),
            "start": start, "end": end}


def pdf_fonts() -> tuple[str, str]:
    regular = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    bold = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if regular.exists() and bold.exists():
        if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVu", str(regular)))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(bold)))
        return "DejaVu", "DejaVu-Bold"
    return "Helvetica", "Helvetica-Bold"


def pdf_bar_chart(items: list[tuple[str, float]], font: str, limit: int = 8) -> Drawing:
    selected = items[:limit]
    height = max(52, len(selected) * 25 + 20)
    drawing = Drawing(748, height)
    if not selected:
        drawing.add(String(8, height - 24, "No tracked time in this period", fontName=font,
                           fontSize=9, fillColor=colors.HexColor("#52677A")))
        return drawing
    maximum = max(value for _, value in selected) or 1
    for idx, (name, value) in enumerate(selected):
        y = height - 28 - idx * 25
        label = name if len(name) <= 28 else name[:25] + "..."
        drawing.add(String(4, y + 2, label, fontName=font, fontSize=9,
                           fillColor=colors.HexColor("#263746")))
        drawing.add(Rect(172, y, 510, 13, fillColor=colors.HexColor("#E6F0F8"),
                         strokeColor=None))
        drawing.add(Rect(172, y, 510 * value / maximum, 13,
                         fillColor=colors.HexColor("#2F6FA3"), strokeColor=None))
        drawing.add(String(695, y + 2, f"{value:,.2f} h", fontName=font, fontSize=9,
                           fillColor=colors.HexColor("#17324F")))
    return drawing


def pdf_table(rows: list[list[object]], widths: list[float], regular: str,
              bold: str, body_style: ParagraphStyle) -> Table:
    head_style = ParagraphStyle("table-head", parent=body_style, fontName=bold,
                                textColor=colors.white)
    converted = []
    for idx, row in enumerate(rows):
        converted.append([Paragraph(escape(str(value)) if value is not None else "",
                                    head_style if idx == 0 else body_style)
                          for value in row])
    table = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F4F8FC")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#D5E4F0")),
    ]))
    return table


def write_pdf(model: dict, path: Path) -> None:
    regular, bold = pdf_fonts()
    page_width, _ = landscape(A4)
    body = ParagraphStyle("body", fontName=regular, fontSize=8.5, leading=11,
                          textColor=colors.HexColor("#263746"))
    title = ParagraphStyle("title", fontName=bold, fontSize=18, leading=24,
                           textColor=colors.HexColor("#17324F"))
    subtitle = ParagraphStyle("subtitle", fontName=regular, fontSize=9, leading=13,
                              textColor=colors.HexColor("#52677A"))
    section = ParagraphStyle("section", fontName=bold, fontSize=11, leading=16,
                             textColor=colors.HexColor("#17324F"), spaceBefore=10,
                             spaceAfter=5)
    temporary = path.with_name(path.stem + ".refreshing.pdf")
    doc = SimpleDocTemplate(str(temporary), pagesize=landscape(A4),
                            leftMargin=42, rightMargin=42, topMargin=30,
                            bottomMargin=25, title="Clockify controlling")
    story = [Paragraph("Clockify time / Jira estimates", title),
             Paragraph(f"{model['start'].isoformat()} to {model['end'].isoformat()}  |  "
                       f"Refreshed {datetime.now().strftime('%Y-%m-%d %H:%M')}", subtitle),
             Spacer(1, 14)]
    kpis = [[Paragraph("Tracked hours", subtitle), Paragraph("People with time", subtitle),
             Paragraph("Tasks", subtitle), Paragraph("Unassigned hours", subtitle)],
            [Paragraph(f"{model['total_hours']:,.2f}", title),
             Paragraph(str(sum(row[1] > 0 for row in model['people'])), title),
             Paragraph(str(sum(bool(row[9]) for row in model['tasks'])), title),
             Paragraph(f"{model['unassigned_hours']:,.2f}", title)]]
    kpi_table = Table(kpis, colWidths=[189] * 4, hAlign="LEFT")
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF5FB")),
        ("TOPPADDING", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
    ]))
    story += [kpi_table, Paragraph("Time by project", section),
              pdf_bar_chart(model["projects"], regular)]
    if len(model["projects"]) > 8:
        story.append(Paragraph("Chart shows the eight projects with the most hours; all projects remain in Excel.", subtitle))
    story.append(Paragraph("Tasks", section))
    story.append(Paragraph("Difference = tracked hours minus Jira original estimate; positive values mean overrun. A dash means no Jira estimate is available.", subtitle))
    task_rows = [["Project", "Task", "Jira original (h)", "Tracked (h)",
                  "Difference (h)", "Top contributor", "Clockify status"]]
    for row in model["tasks"]:
        estimate = row[2]
        task_rows.append([row[0], row[1], f"{estimate:,.2f}" if estimate is not None else "-",
                          f"{row[3]:,.2f}", f"{row[3] - estimate:+,.2f}" if estimate is not None else "-",
                          row[7] or "-", row[5] or "-"])
    story.append(pdf_table(task_rows, [95, 220, 85, 72, 82, 127, 75], regular, bold, body))
    story += [PageBreak(), Paragraph("Time by person", title),
              Paragraph(f"{model['start'].isoformat()} to {model['end'].isoformat()}", subtitle),
              Spacer(1, 12),
              pdf_bar_chart([(row[0], row[1]) for row in model["people"]], regular)]
    if len(model["people"]) > 8:
        story.append(Paragraph("Chart shows the eight people with the most hours; the table includes everyone.", subtitle))
    story.append(Paragraph("Person totals", section))
    person_rows = [["Person", "Tracked (h)", "Billable (h)", "Non-billable (h)",
                    "Entries", "Projects"]]
    for row in model["people"]:
        person_rows.append([row[0], f"{row[1]:,.2f}", f"{row[2]:,.2f}",
                            f"{row[3]:,.2f}", row[5], row[6]])
    story.append(pdf_table(person_rows, [181, 110, 110, 125, 100, 130], regular, bold, body))
    story.append(Paragraph("Person and task", section))
    detail_rows = [["Person", "Project", "Task", "Hours"]]
    for person, project, task, value in model["person_tasks"]:
        detail_rows.append([person, project, task, f"{value:,.2f}"])
    story.append(pdf_table(detail_rows, [188, 178, 290, 100], regular, bold, body))
    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D5E4F0"))
        canvas.line(42, 28, page_width - 42, 28)
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#52677A"))
        canvas.drawString(42, 17, "Source: Clockify time entries and Jira original estimates")
        canvas.drawRightString(page_width - 42, 17, f"Page {document.page}")
        canvas.restoreState()

    try:
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description="Refresh Clockify hours and Jira original estimates")
    parser.add_argument("--workbook", type=Path, default=HERE / "Clockify_Controlling.xlsx")
    parser.add_argument("--pdf", type=Path, help="PDF output (default: workbook name with .pdf)")
    parser.add_argument("--env-file", type=Path, default=HERE / ".env",
                        help="Private Clockify and Jira configuration (default: .env beside script)")
    parser.add_argument("--since", required=True, type=date.fromisoformat, help="First work date, YYYY-MM-DD")
    parser.add_argument("--until", type=date.fromisoformat, default=date.today(), help="Last work date, YYYY-MM-DD")
    args = parser.parse_args()
    load_env(args.env_file)
    key, workspace = os.getenv("CLOCKIFY_API_KEY"), os.getenv("CLOCKIFY_WORKSPACE_ID")
    if not key or not workspace:
        raise RuntimeError("Set CLOCKIFY_API_KEY and CLOCKIFY_WORKSPACE_ID in .env or environment")
    jira_url = os.getenv("JIRA_URL")
    jira_email = os.getenv("JIRA_EMAIL")
    jira_token = os.getenv("JIRA_API_TOKEN")
    jira_project = os.getenv("JIRA_PROJECT_KEY")
    missing_jira = [name for name, value in (("JIRA_URL", jira_url), ("JIRA_EMAIL", jira_email),
                                             ("JIRA_API_TOKEN", jira_token)) if not value]
    if missing_jira:
        raise RuntimeError(f"Set {', '.join(missing_jira)} in .env or environment")
    if args.until < args.since:
        raise RuntimeError("--until precedes --since")
    api_base = os.getenv("CLOCKIFY_API_BASE", "https://api.clockify.me/api/v1")
    report_base = os.getenv("CLOCKIFY_REPORT_BASE", "https://reports.api.clockify.me/v1")
    projects = pages(api_base, f"/workspaces/{workspace}/projects", key)
    tasks = []
    for project in projects:
        tasks.extend(pages(api_base, f"/workspaces/{workspace}/projects/{project['id']}/tasks", key))
    users = pages(api_base, f"/workspaces/{workspace}/users", key)
    entries = report_entries(report_base, workspace, key, args.since, args.until)
    issues = jira_issues(jira_url, jira_email, jira_token, jira_project)
    model = refresh(args.workbook, projects, tasks, users, entries, issues, args.since, args.until)
    pdf_path = args.pdf or args.workbook.with_suffix(".pdf")
    write_pdf(model, pdf_path)
    print(f"Updated {args.workbook} and {pdf_path}: "
          f"{len(model['tasks'])} task rows, {len(model['entries'])} time entries, "
          f"{len(model['people'])} people, "
          f"{sum(row[12] == 'No Jira match' for row in model['tasks'])} unmatched, "
          f"{sum(row[12].startswith('Ambiguous:') for row in model['tasks'])} ambiguous")


if __name__ == "__main__":
    try:
        main()
    except (RuntimeError, KeyError, ValueError, OSError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1)
